import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError

from operations.models import InventoryTransaction, StorageLocation, Supplier, Warehouse
from operations.services import record_inventory_change

from .models import (
    Category,
    DosageForm,
    InventoryBatch,
    Manufacturer,
    Medicine,
    MedicineImportJob,
    RecordStatus,
)

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_ROWS_PER_SHEET = 5000

MEDICINE_HEADERS = (
    "Medicine Code / SKU",
    "Medicine Name",
    "Generic Name / Active Ingredient",
    "Dari Name",
    "Strength",
    "Dosage Form",
    "Category",
    "Manufacturer",
    "Country of Manufacture",
    "Pack Size",
    "Prescription Required?",
    "Selling Price (AFN)",
    "Purchase/Cost Price (AFN)",
    "Barcode",
    "Description",
    "Storage Instructions",
    "Active?",
)

STOCK_HEADERS = (
    "Medicine Code / SKU",
    "Batch Number",
    "Quantity in Stock",
    "Expiry Date",
    "Purchase Price (AFN)",
    "Selling Price (AFN)",
    "Warehouse/Location",
    "Rack/Shelf",
    "Supplier",
    "Received Date",
)


def import_template_bytes():
    workbook = Workbook()
    medicines = workbook.active
    medicines.title = "Medicines"
    medicines.append(MEDICINE_HEADERS)
    stock = workbook.create_sheet("Stock - Batches")
    stock.append(STOCK_HEADERS)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _sheet_rows(sheet, expected_headers):
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ValidationError({"file": f"Sheet '{sheet.title}' is empty."}) from exc
    headers = tuple(_clean(value) for value in raw_headers)
    if headers != expected_headers:
        raise ValidationError(
            {"file": f"Sheet '{sheet.title}' headers must exactly match the import template."}
        )
    result = []
    for row_number, values in enumerate(rows, start=2):
        if row_number > MAX_ROWS_PER_SHEET + 1:
            raise ValidationError({"file": f"Sheet '{sheet.title}' exceeds 5,000 data rows."})
        if any(isinstance(value, str) and value.lstrip().startswith("=") for value in values):
            raise ValidationError(
                {"file": f"Formulas are not allowed ({sheet.title} row {row_number})."}
            )
        normalized = {header: _clean(value) for header, value in zip(headers, values, strict=False)}
        if any(normalized.values()):
            result.append((row_number, normalized))
    return result


def _read_upload(upload):
    if upload.size > MAX_IMPORT_BYTES:
        raise ValidationError({"file": "Import files must be 10 MB or smaller."})
    suffix = Path(upload.name).suffix.lower()
    content = upload.read()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValidationError({"file": "CSV files must use UTF-8 encoding."}) from exc
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            raise ValidationError({"file": "The CSV file is empty."})
        if tuple(_clean(value) for value in rows[0]) != MEDICINE_HEADERS:
            raise ValidationError({"file": "CSV headers must match the Medicines template."})
        if len(rows) - 1 > MAX_ROWS_PER_SHEET:
            raise ValidationError({"file": "CSV exceeds 5,000 data rows."})
        medicines = []
        for number, values in enumerate(rows[1:], start=2):
            normalized = {
                header: _clean(value)
                for header, value in zip(MEDICINE_HEADERS, values, strict=False)
            }
            if any(normalized.values()):
                medicines.append((number, normalized))
        return medicines, []
    if suffix != ".xlsx":
        raise ValidationError({"file": "Upload an XLSX workbook or UTF-8 CSV file."})
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValidationError({"file": "The XLSX workbook is malformed or unreadable."}) from exc
    required = {"Medicines", "Stock - Batches"}
    if not required.issubset(workbook.sheetnames):
        raise ValidationError({"file": "XLSX must contain 'Medicines' and 'Stock - Batches'."})
    try:
        return (
            _sheet_rows(workbook["Medicines"], MEDICINE_HEADERS),
            _sheet_rows(workbook["Stock - Batches"], STOCK_HEADERS),
        )
    finally:
        workbook.close()


def _boolean(value, field, errors, sheet, row):
    normalized = value.casefold()
    if normalized in {"yes", "y", "true", "1"}:
        return True
    if normalized in {"no", "n", "false", "0"}:
        return False
    errors.append(_error(sheet, row, field, "Use Yes or No."))
    return False


def _decimal(value, field, errors, sheet, row, required=True):
    if not value and not required:
        return None
    try:
        result = Decimal(value)
    except (InvalidOperation, ValueError):
        errors.append(_error(sheet, row, field, "Enter a valid non-negative amount."))
        return None
    if not result.is_finite() or result < 0:
        errors.append(_error(sheet, row, field, "Enter a valid non-negative amount."))
        return None
    return str(result.quantize(Decimal("0.01")))


def _integer(value, field, errors, sheet, row):
    try:
        result = int(value)
    except (TypeError, ValueError):
        errors.append(_error(sheet, row, field, "Enter a whole number greater than zero."))
        return None
    if result <= 0:
        errors.append(_error(sheet, row, field, "Enter a whole number greater than zero."))
        return None
    return result


def _date(value, field, errors, sheet, row, *, future=False):
    try:
        result = date.fromisoformat(value)
    except (TypeError, ValueError):
        errors.append(_error(sheet, row, field, "Use an ISO date in YYYY-MM-DD format."))
        return None
    if future and result <= timezone.localdate():
        errors.append(_error(sheet, row, field, "Expiry date must be in the future."))
        return None
    return result.isoformat()


def _error(sheet, row, field, message):
    return {"sheet": sheet, "row": row, "field": field, "message": message}


def _resolve_location(value, rack_shelf):
    query = StorageLocation.objects.select_related("warehouse__branch").filter(is_active=True)
    direct = list(query.filter(code__iexact=value)[:2])
    if len(direct) == 1:
        return direct[0]
    warehouse_code, separator, location_code = value.partition("/")
    if separator and location_code:
        return query.filter(
            warehouse__code__iexact=warehouse_code.strip(), code__iexact=location_code.strip()
        ).first()
    warehouse = Warehouse.objects.filter(code__iexact=value, is_active=True).first()
    if not warehouse:
        return None
    rack, separator, shelf = rack_shelf.partition("/")
    if not separator:
        return None
    return query.filter(
        warehouse=warehouse, rack__iexact=rack.strip(), shelf__iexact=shelf.strip()
    ).first()


def build_import_preview(*, upload, mode, actor):
    if mode not in MedicineImportJob.Mode.values:
        raise ValidationError({"mode": "Select a supported import mode."})
    medicine_rows, stock_rows = _read_upload(upload)
    errors = []
    medicines = []
    batches = []
    seen_skus = set()
    seen_barcodes = set()
    known_skus = set(Medicine.objects.values_list("sku", flat=True))
    barcode_owners = dict(
        Medicine.objects.exclude(barcode__isnull=True).values_list("barcode", "sku")
    )

    for row, values in medicine_rows:
        sku = values[MEDICINE_HEADERS[0]].upper()
        name = values[MEDICINE_HEADERS[1]]
        dosage_name = values[MEDICINE_HEADERS[5]]
        barcode = values[MEDICINE_HEADERS[13]] or None
        if not sku:
            errors.append(_error("Medicines", row, MEDICINE_HEADERS[0], "SKU is required."))
        elif sku in seen_skus:
            errors.append(_error("Medicines", row, MEDICINE_HEADERS[0], "Duplicate SKU in file."))
        seen_skus.add(sku)
        if not name:
            errors.append(
                _error("Medicines", row, MEDICINE_HEADERS[1], "Medicine name is required.")
            )
        dosage = (
            DosageForm.objects.filter(name__iexact=dosage_name).first()
            or DosageForm.objects.filter(code__iexact=dosage_name).first()
        )
        if not dosage:
            errors.append(
                _error("Medicines", row, MEDICINE_HEADERS[5], "Dosage form was not found.")
            )
        category_name = values[MEDICINE_HEADERS[6]]
        category = (
            Category.objects.filter(name__iexact=category_name).first()
            or Category.objects.filter(name_dari__iexact=category_name).first()
        )
        if category_name and not category:
            errors.append(_error("Medicines", row, MEDICINE_HEADERS[6], "Category was not found."))
        manufacturer_name = values[MEDICINE_HEADERS[7]]
        manufacturer = Manufacturer.objects.filter(name__iexact=manufacturer_name).first()
        country = values[MEDICINE_HEADERS[8]].upper()
        if country and len(country) != 2:
            errors.append(
                _error("Medicines", row, MEDICINE_HEADERS[8], "Use a two-letter country code.")
            )
        if not manufacturer_name:
            errors.append(
                _error("Medicines", row, MEDICINE_HEADERS[7], "Manufacturer is required.")
            )
        if barcode and (
            barcode in seen_barcodes
            or (barcode in barcode_owners and barcode_owners[barcode] != sku)
        ):
            errors.append(_error("Medicines", row, MEDICINE_HEADERS[13], "Barcode is duplicated."))
        if barcode:
            seen_barcodes.add(barcode)
        exists = sku in known_skus
        action = "CREATE"
        if exists:
            if mode == MedicineImportJob.Mode.CREATE_NEW:
                errors.append(
                    _error(
                        "Medicines",
                        row,
                        MEDICINE_HEADERS[0],
                        "SKU already exists; select update or skip mode.",
                    )
                )
                action = "SKIP"
            else:
                action = "UPDATE" if mode == MedicineImportJob.Mode.UPDATE_EXISTING else "SKIP"
        medicines.append(
            {
                "row": row,
                "action": action,
                "sku": sku,
                "brand_name": name,
                "generic_name": values[MEDICINE_HEADERS[2]],
                "brand_name_dari": values[MEDICINE_HEADERS[3]],
                "strength": values[MEDICINE_HEADERS[4]],
                "dosage_form_id": str(dosage.id) if dosage else None,
                "category_id": str(category.id) if category else None,
                "manufacturer_id": str(manufacturer.id) if manufacturer else None,
                "manufacturer_name": manufacturer_name,
                "country_code": country,
                "pack_size": values[MEDICINE_HEADERS[9]],
                "prescription_required": _boolean(
                    values[MEDICINE_HEADERS[10]], MEDICINE_HEADERS[10], errors, "Medicines", row
                ),
                "default_selling_price": _decimal(
                    values[MEDICINE_HEADERS[11]],
                    MEDICINE_HEADERS[11],
                    errors,
                    "Medicines",
                    row,
                    required=False,
                ),
                "default_purchase_price": _decimal(
                    values[MEDICINE_HEADERS[12]],
                    MEDICINE_HEADERS[12],
                    errors,
                    "Medicines",
                    row,
                    required=False,
                ),
                "barcode": barcode,
                "description": values[MEDICINE_HEADERS[14]],
                "storage_instructions": values[MEDICINE_HEADERS[15]],
                "status": Medicine.Status.ACTIVE
                if _boolean(
                    values[MEDICINE_HEADERS[16]], MEDICINE_HEADERS[16], errors, "Medicines", row
                )
                else Medicine.Status.INACTIVE,
            }
        )

    usable_skus = known_skus | seen_skus
    seen_batches = set()
    medicine_defaults = {item["sku"]: item for item in medicines}
    for row, values in stock_rows:
        sku = values[STOCK_HEADERS[0]].upper()
        batch_number = values[STOCK_HEADERS[1]]
        if sku not in usable_skus:
            errors.append(
                _error(
                    "Stock - Batches",
                    row,
                    STOCK_HEADERS[0],
                    "Medicine SKU was not found in the database or Medicines sheet.",
                )
            )
        if not batch_number:
            errors.append(
                _error("Stock - Batches", row, STOCK_HEADERS[1], "Batch number is required.")
            )
        location = _resolve_location(values[STOCK_HEADERS[6]], values[STOCK_HEADERS[7]])
        if not location:
            errors.append(
                _error(
                    "Stock - Batches",
                    row,
                    STOCK_HEADERS[6],
                    "Active warehouse/storage location was not found.",
                )
            )
        supplier_value = values[STOCK_HEADERS[8]]
        supplier = (
            Supplier.objects.filter(code__iexact=supplier_value, is_active=True).first()
            or Supplier.objects.filter(company_name__iexact=supplier_value, is_active=True).first()
        )
        if not supplier:
            errors.append(
                _error("Stock - Batches", row, STOCK_HEADERS[8], "Active supplier was not found.")
            )
        key = (sku, batch_number, str(location.warehouse.branch_id) if location else "")
        if key in seen_batches:
            errors.append(
                _error("Stock - Batches", row, STOCK_HEADERS[1], "Duplicate batch in workbook.")
            )
        seen_batches.add(key)
        existing_batch = False
        if location and sku and batch_number:
            existing_batch = InventoryBatch.objects.filter(
                branch=location.warehouse.branch, medicine__sku=sku, batch_number=batch_number
            ).exists()
        default = medicine_defaults.get(sku, {})
        batches.append(
            {
                "row": row,
                "action": "SKIP" if existing_batch else "CREATE",
                "sku": sku,
                "batch_number": batch_number,
                "quantity": _integer(
                    values[STOCK_HEADERS[2]], STOCK_HEADERS[2], errors, "Stock - Batches", row
                ),
                "expiry_date": _date(
                    values[STOCK_HEADERS[3]],
                    STOCK_HEADERS[3],
                    errors,
                    "Stock - Batches",
                    row,
                    future=True,
                ),
                "purchase_price": _decimal(
                    values[STOCK_HEADERS[4]] or default.get("default_purchase_price", ""),
                    STOCK_HEADERS[4],
                    errors,
                    "Stock - Batches",
                    row,
                ),
                "selling_price": _decimal(
                    values[STOCK_HEADERS[5]] or default.get("default_selling_price", ""),
                    STOCK_HEADERS[5],
                    errors,
                    "Stock - Batches",
                    row,
                ),
                "storage_location_id": str(location.id) if location else None,
                "supplier_id": str(supplier.id) if supplier else None,
                "received_date": _date(
                    values[STOCK_HEADERS[9]], STOCK_HEADERS[9], errors, "Stock - Batches", row
                ),
            }
        )

    summary = {
        "medicine_rows": len(medicines),
        "new_medicines": sum(item["action"] == "CREATE" for item in medicines),
        "updated_medicines": sum(item["action"] == "UPDATE" for item in medicines),
        "skipped_medicines": sum(item["action"] == "SKIP" for item in medicines),
        "batch_rows": len(batches),
        "new_batches": sum(item["action"] == "CREATE" for item in batches),
        "duplicate_batches": sum(item["action"] == "SKIP" for item in batches),
        "error_count": len(errors),
    }
    return MedicineImportJob.objects.create(
        uploaded_by=actor,
        original_filename=Path(upload.name).name[:255],
        mode=mode,
        normalized_payload={"medicines": medicines, "batches": batches},
        validation_report={"summary": summary, "errors": errors},
    )


@transaction.atomic
def confirm_import(*, job, actor):
    locked = MedicineImportJob.objects.select_for_update().get(pk=job.pk)
    if locked.status != MedicineImportJob.Status.PREVIEWED:
        raise ValidationError({"job": "This import has already been confirmed."})
    if locked.validation_report.get("errors"):
        raise ValidationError({"job": "Fix every preview error before confirming the import."})
    created = updated = skipped = batches_created = batches_skipped = 0
    for item in locked.normalized_payload.get("medicines", []):
        if item["action"] == "SKIP":
            skipped += 1
            continue
        manufacturer = Manufacturer.objects.filter(pk=item["manufacturer_id"]).first()
        if not manufacturer:
            manufacturer = Manufacturer.objects.filter(
                name__iexact=item["manufacturer_name"]
            ).first()
        if not manufacturer:
            manufacturer = Manufacturer.objects.create(
                name=item["manufacturer_name"],
                country_code=item["country_code"],
                status=RecordStatus.ACTIVE,
            )
        fields = {
            "brand_name": item["brand_name"],
            "brand_name_dari": item["brand_name_dari"],
            "generic_name": item["generic_name"],
            "manufacturer": manufacturer,
            "category_id": item["category_id"],
            "dosage_form_id": item["dosage_form_id"],
            "strength": item["strength"],
            "pack_size": item["pack_size"],
            "barcode": item["barcode"],
            "description": item["description"],
            "storage_instructions": item["storage_instructions"],
            "prescription_required": item["prescription_required"],
            "status": item["status"],
        }
        if item["action"] == "CREATE":
            medicine = Medicine(sku=item["sku"], **fields)
            try:
                medicine.full_clean()
            except DjangoValidationError as exc:
                raise ValidationError(exc.message_dict) from exc
            medicine.save()
            created += 1
        else:
            medicine = Medicine.objects.select_for_update().get(sku=item["sku"])
            for field, value in fields.items():
                setattr(medicine, field, value)
            try:
                medicine.full_clean()
            except DjangoValidationError as exc:
                raise ValidationError(exc.message_dict) from exc
            medicine.save()
            updated += 1
    for item in locked.normalized_payload.get("batches", []):
        if item["action"] == "SKIP":
            batches_skipped += 1
            continue
        medicine = Medicine.objects.get(sku=item["sku"])
        location = StorageLocation.objects.select_related("warehouse__branch").get(
            pk=item["storage_location_id"]
        )
        batch = InventoryBatch.objects.create(
            medicine=medicine,
            branch=location.warehouse.branch,
            storage_location=location,
            batch_number=item["batch_number"],
            expiry_date=item["expiry_date"],
            available_quantity=0,
            purchase_price=item["purchase_price"],
            selling_price=item["selling_price"],
            quality_status=InventoryBatch.QualityStatus.PENDING,
            is_active=False,
        )
        supplier = Supplier.objects.get(pk=item["supplier_id"])
        record_inventory_change(
            batch=batch,
            transaction_type=InventoryTransaction.TransactionType.PURCHASE,
            direction=InventoryTransaction.Direction.ADDITION,
            quantity=item["quantity"],
            actor=actor,
            reference=f"IMPORT-{locked.id}",
            unit_cost=item["purchase_price"],
            notes=f"Bulk import; supplier {supplier.code}; received {item['received_date']}",
        )
        batches_created += 1
    result = {
        "created_medicines": created,
        "updated_medicines": updated,
        "skipped_medicines": skipped,
        "created_batches": batches_created,
        "skipped_batches": batches_skipped,
    }
    locked.status = MedicineImportJob.Status.COMPLETED
    locked.confirmed_at = timezone.now()
    locked.result_report = result
    locked.save(update_fields=("status", "confirmed_at", "result_report", "updated_at"))
    return locked
